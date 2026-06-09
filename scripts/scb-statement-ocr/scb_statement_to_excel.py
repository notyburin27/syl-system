#!/usr/bin/env python3
"""Convert SCANNED SCB savings-account statements (image-only PDFs) to Excel.

The in-app converter (lib/utils/statementPdfParser.ts) only handles digital PDFs
with a text layer. Scanned statements (e.g. from a photocopier) have no text, and
the scanned "STATEMENT OF SAVING ACCOUNT" layout differs from the formats that
parser knows. This standalone script handles that case via OCR.

How it works
------------
1. Render each page to an image (pdftoppm, 300 dpi).
2. OCR each page to word boxes (tesseract, tha+eng, TSV).
3. Slot words into columns by their x-position (the table is right-aligned, so the
   right edge of each money column is stable): debit ~rx 957, credit ~rx 1248,
   balance ~rx 1610 at 300 dpi.
4. Reconcile every row against the printed running balance. Because each balance =
   previous balance +/- the transaction amount, the balance column acts as a
   checksum: it fixes most OCR digit slips automatically and flags the rest.

OCR is imperfect on Thai scans, so the result is NOT guaranteed correct. The
balance chain is reconciled end to end, and any row it could not verify is marked
in the "Check" column and duplicated into a "Review" sheet for manual checking
against the PDF. Thai descriptions/names come out garbled; numbers are the focus.

Requirements
------------
    brew install poppler tesseract tesseract-lang     # pdftoppm + tesseract(+tha)
    pip3 install --user openpyxl

Usage
-----
    python3 scb_statement_to_excel.py OUT.xlsx IN1.pdf [IN2.pdf ...]
    python3 scb_statement_to_excel.py OUT.xlsx --max-balance 15000000 IN.pdf

PDFs are processed in the order given; pass split scan sets in chronological order
so the running balance stays continuous. Non-statement pages (receipts, blank
covers) yield no rows and are skipped automatically.
"""
import csv, re, sys, os, glob, tempfile, subprocess, shutil
from collections import Counter

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl not installed. Run: pip3 install --user openpyxl")

DPI = 300
EPS = 0.005
# Account-specific guard: balances/amounts larger than this are treated as OCR
# digit-insertions (extra digits) rather than real values. Tune with --max-balance
# to a little above the account's real peak balance. 0 disables the guard.
MAX_BAL = 15_000_000

DATE = re.compile(r'\d{2}/\d{2}/\d{2}')
TIMEPAT = re.compile(r'(\d{1,2})[:.;](\d{2})')
CORE = re.compile(r'\d[\d.,]*\d|\d')
CHAN = {'ENET', 'BCMS', 'TELL', 'ATM', 'SYSG', 'MNET', 'BNET', 'CASH'}
KNOWN_CODES = {'X1', 'X2', 'XW', 'XO', 'FE', 'C1', 'C2', 'C0', 'CO', 'CW', 'QN',
               'IN', 'CX', 'DR', 'CR', 'IB', 'PB', 'ON'}
DIGITMAP = str.maketrans({'O': '0', 'o': '0', 'Q': '0', 'D': '0', 'l': '1',
                          'I': '1', '§': '5', 'S': '5', 'B': '8', 'U': '0'})
BF_KW = ('BROUGHT', 'FORWARD', 'ยกมา')
CLEAN = {'OK', 'FIX_AMT', 'FIX_BAL', 'DERIVED_AMT'}
REVIEW_FLAGS = {'REVIEW', 'BF_MISMATCH', 'NO_BAL', 'NO_ANCHOR'}
HEADERS = ['Date', 'Time', 'Code', 'Channel', 'Debit/Credit', 'Balance/Baht',
           'Description/Note', 'Description_Clean', 'Name', 'Note', 'Page', 'Check']


# ----------------------------------------------------------------- render + OCR
def require_tool(name):
    if shutil.which(name) is None:
        sys.exit(f"'{name}' not found. Install with: brew install poppler tesseract tesseract-lang")


def page_count(pdf):
    out = subprocess.run(['pdfinfo', pdf], capture_output=True, text=True).stdout
    m = re.search(r'Pages:\s+(\d+)', out)
    return int(m.group(1)) if m else 0


def ocr_page(pdf, page, workdir):
    """Render one PDF page and OCR it; return the path to a tesseract TSV file."""
    stem = os.path.join(workdir, 'page')
    subprocess.run(['pdftoppm', '-png', '-r', str(DPI), '-f', str(page), '-l', str(page),
                    pdf, stem], check=True, capture_output=True)
    produced = sorted(glob.glob(stem + '-*.png'))
    if not produced:
        return None
    png = produced[0]
    tsv_base = os.path.join(workdir, 'page')
    subprocess.run(['tesseract', png, tsv_base, '-l', 'tha+eng', '--psm', '6', 'tsv'],
                   check=True, capture_output=True)
    os.remove(png)
    return tsv_base + '.tsv'


# ------------------------------------------------------------------ extraction
def load_words(path):
    out = []
    with open(path) as f:
        for r in csv.DictReader(f, delimiter='\t', quoting=csv.QUOTE_NONE):
            if r.get('level') != '5':
                continue
            t = (r['text'] or '').strip()
            if not t:
                continue
            x, y, w, h = int(r['left']), int(r['top']), int(r['width']), int(r['height'])
            out.append(dict(x=x, y=y, w=w, h=h, rx=x + w, cx=x + w / 2, cy=y + h / 2, text=t))
    return out


def cluster_rows(words, ytol=14):
    ws = sorted(words, key=lambda d: d['cy'])
    rows, cur, cy = [], [], None
    for w in ws:
        if cy is None or abs(w['cy'] - cy) <= ytol:
            cur.append(w); cy = w['cy'] if cy is None else cy * 0.6 + w['cy'] * 0.4
        else:
            rows.append(cur); cur = [w]; cy = w['cy']
    if cur:
        rows.append(cur)
    return rows


def is_num_frag(s):
    if any(c.isdigit() for c in s):
        return True
    return bool(re.match(r'^[.,][OoQ]{1,3}', s))


def is_money_like(s):
    return bool(re.search(r'\d', s)) and ((',' in s) or ('.' in s) or len(re.sub(r'\D', '', s)) >= 3)


def parse_amount(tokens):
    """Merge band fragments (x order) -> value. The bank always prints exactly 2
    decimals, so value = (all digits of the numeric core) / 100. Robust to lost
    decimal points ('674.241.00', '1,021,24700'); only fixes an extra 3rd decimal."""
    if not tokens:
        return None
    toks = sorted(tokens, key=lambda d: d['x'])
    raw = ''.join(t['text'].translate(DIGITMAP) for t in toks)
    m = CORE.search(raw)
    if not m:
        return None
    core = m.group(0)
    digits = re.sub(r'\D', '', core)
    em = re.search(r'\.(\d{3,})$', core)
    if em and len(digits) > 2:
        digits = digits[:-(len(em.group(1)) - 2)]
    return int(digits) / 100.0 if digits else None


def clean_code(tok):
    s = tok.upper().replace('%', 'X').replace('*', 'X')
    s = re.sub(r'[^A-Z0-9]', '', s)
    if s in KNOWN_CODES:
        return s
    return s[-2:] if len(s) > 2 else s


def clean_time(tok):
    m = TIMEPAT.search(tok.translate(DIGITMAP))
    return f"{int(m.group(1)):02d}:{m.group(2)}" if m else ''


def join_desc(tokens):
    toks = sorted(tokens, key=lambda d: d['x'])
    parts, prev_ascii = [], None
    for t in toks:
        txt = t['text']; a = txt.isascii()
        if parts and prev_ascii is False and not a:
            parts[-1] += txt
        else:
            parts.append(txt)
        prev_ascii = a
    return ' '.join(parts).strip()


def extract_name(desc):
    m = re.search(r'(นาย|นางสาว|นาง|น\.ส\.|MR\.?|MS\.?|MRS\.?)\s*([ก-๙A-Za-z\. ]+)', desc)
    if m:
        return (m.group(1) + ' ' + m.group(2)).strip()[:40]
    m = re.search(r'(OWN|TO|FROM)\s+([A-Z][A-Z \.]{2,})', desc)
    return m.group(2).strip()[:40] if m else ''


def extract_page(tsv_path, page_tag):
    rows = cluster_rows(load_words(tsv_path))
    recs = []
    for r in rows:
        toks = sorted(r, key=lambda d: d['x'])
        linetext = ' '.join(t['text'] for t in toks).upper()
        date = time = code = chan = None
        debit_t, credit_t, bal_t, desc_t, date_band = [], [], [], [], []
        for w in toks:
            if w['text'] in CHAN and chan is None:
                chan = w['text']; continue
            if w['rx'] < 300:
                date_band.append(w)
                if date is None and DATE.search(w['text']):
                    date = DATE.search(w['text']).group(0)
                continue
        for w in toks:
            t = w['text']
            if date and DATE.search(t) and w['rx'] < 300:
                continue
            if t in CHAN:
                continue
            cx, rx = w['cx'], w['rx']
            if time is None and 285 < rx < 412 and TIMEPAT.search(t.translate(DIGITMAP)):
                time = clean_time(t); continue
            if code is None and 395 < cx < 515 and not is_money_like(t) and re.search(r'[A-Za-z%*0-9]', t):
                cc = clean_code(t)
                if cc:
                    code = cc; continue
            if is_num_frag(t):
                if 895 < rx < 1005:
                    debit_t.append(w); continue
                if 1175 < rx < 1310:
                    credit_t.append(w); continue
                if 1340 < cx < 1648:
                    bal_t.append(w); continue
            if cx >= 1648:
                desc_t.append(w)
        if date is None and date_band:
            dd = ''.join(re.sub(r'\D', '', t['text'].translate(DIGITMAP))
                         for t in sorted(date_band, key=lambda d: d['x']))
            if len(dd) == 6:
                date = f"{dd[0:2]}/{dd[2:4]}/{dd[4:6]}"
        debit, credit, balance = parse_amount(debit_t), parse_amount(credit_t), parse_amount(bal_t)
        desc = join_desc(desc_t)
        has_amt = debit is not None or credit is not None
        is_bf = (date is None) and balance is not None and not has_amt and any(k in linetext for k in BF_KW)
        is_data = has_amt and balance is not None and (chan is not None or code is not None or time)
        if is_bf:
            recs.append(dict(page=page_tag, date='', time='', code='', chan='', debit=None,
                             credit=None, balance=balance, desc='BALANCE BROUGHT FORWARD', bf=True))
        elif is_data:
            recs.append(dict(page=page_tag, date=date or '', time=time or '', code=code or '',
                             chan=chan or '', debit=debit, credit=credit, balance=balance,
                             desc=desc, bf=False))
    return recs


# --------------------------------------------------------------- reconciliation
def signed_change(rec):
    if rec is None:
        return None
    d, c = rec['debit'], rec['credit']
    if d is not None and c is not None:
        return -d if d >= c else c
    if d is not None:
        return -d
    if c is not None:
        return c
    return None


def cents(v):
    return None if v is None else int(round(v * 100))


def edit_dist(a, b):
    if a is None or b is None:
        return 99
    sa, sb = str(abs(a)), str(abs(b))
    m, n = len(sa), len(sb)
    dp = list(range(n + 1))
    for i in range(1, m + 1):
        prev = dp[0]; dp[0] = i
        for j in range(1, n + 1):
            cur = dp[j]
            dp[j] = min(dp[j] + 1, dp[j - 1] + 1, prev + (sa[i - 1] != sb[j - 1]))
            prev = cur
    return dp[n]


def reconcile(recs, max_bal):
    running = None; out = []; n = len(recs); prev_clean = False
    for i, rec in enumerate(recs):
        r2 = dict(rec)
        if rec['bf']:
            b = rec['balance']; st = 'BF'
            if running is not None and b is not None and abs(b - running) > EPS:
                if prev_clean and edit_dist(cents(b), cents(running)) <= 1:
                    st = 'BF_FIX'                      # keep running (continuity)
                else:
                    running = b; st = 'BF_MISMATCH'
            elif b is not None:
                running = b
            r2.update(signed=None, status=st, final_balance=running); out.append(r2); continue

        amt = signed_change(rec); b = rec['balance']; old = running
        suspect = False
        if max_bal and b is not None and b > max_bal:
            b = None; suspect = True
        if max_bal and amt is not None and abs(amt) > max_bal:
            amt = None; suspect = True
        if suspect and amt is None and b is None and running is not None:
            r2.update(signed=0.0, status='REVIEW', final_balance=running)
            out.append(r2); prev_clean = False; continue

        if running is None:
            running = b if b is not None else (amt or 0.0)
            r2.update(signed=amt, status='NO_ANCHOR', final_balance=running); out.append(r2); continue
        exp = None if amt is None else old + amt
        if amt is not None and b is not None and abs(b - exp) < EPS:
            running = b; r2.update(signed=amt, status='OK')
        elif amt is None and b is not None:
            r2.update(signed=b - old, status='DERIVED_AMT'); running = b
        elif b is None and amt is not None:
            running = exp; r2.update(signed=amt, status='NO_BAL')
        else:
            # printed balance (b) disagrees with chain estimate (exp). Confirm each
            # hypothesis by following OCR amounts forward up to 3 rows to a printed balance.
            B1, B2 = exp, b

            def confirmed(start):
                acc = start
                for k in range(1, 4):
                    j = i + k
                    if j >= n or recs[j]['bf']:
                        break
                    a = signed_change(recs[j])
                    if a is None:
                        break
                    acc += a; nbk = recs[j]['balance']
                    if nbk is not None and abs(nbk - acc) < EPS:
                        return True
                return False

            b1_conf, b2_conf = confirmed(B1), confirmed(B2)
            if b1_conf and not b2_conf:
                running = B1; r2.update(signed=amt, status='FIX_BAL')        # printed balance misread
            elif b2_conf and not b1_conf:
                running = B2; r2.update(signed=b - old, status='FIX_AMT')    # amount misread
            else:
                running = b; r2.update(signed=b - old, status='REVIEW')      # trust printed balance, flag
        if suspect and r2['status'] not in REVIEW_FLAGS:
            r2['status'] = 'REVIEW'
        r2['final_balance'] = running; out.append(r2)
        prev_clean = r2['status'] in CLEAN
    return out


def sanitize_dates(rec):
    """Fix OCR-mangled dd/mm/yy: keep plausible dates, else carry forward the last good one.
    Year window is intentionally wide (15-29) so it works across statements."""
    prev = ''
    for r in rec:
        if r['bf']:
            continue
        m = re.match(r'(\d{2})/(\d{2})/(\d{2})$', r['date'] or '')
        ok = bool(m) and 1 <= int(m.group(1)) <= 31 and 1 <= int(m.group(2)) <= 12 and 15 <= int(m.group(3)) <= 29
        if ok:
            prev = r['date']
        else:
            r['date'] = prev


# ------------------------------------------------------------------------ excel
def _style_header(ws):
    hf = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='4472C4')
    for c in ws[1]:
        c.font = hf; c.fill = fill; c.alignment = Alignment(horizontal='center', vertical='center')
    for i, wd in enumerate([11, 7, 7, 9, 15, 15, 46, 40, 28, 12, 6, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = 'A2'


def _row_for(r):
    if r['bf']:
        return ['', '', '', '', '', r['final_balance'], r['desc'], '', '', '', r['page'], r['status']]
    return [r['date'], r['time'], r['code'], r['chan'], r['signed'], r['final_balance'],
            r['desc'], r['desc'], extract_name(r['desc']), '', r['page'], r['status']]


def write_excel(rec, path):
    red = Font(color='FF0000'); green = Font(color='008000'); orange = Font(color='C55A11')
    wb = Workbook(); ws = wb.active; ws.title = 'Statement'
    ws.append(HEADERS); _style_header(ws)
    for r in rec:
        ws.append(_row_for(r)); rr = ws.max_row
        ws.cell(rr, 5).number_format = '#,##0.00'; ws.cell(rr, 6).number_format = '#,##0.00'
        if not r['bf'] and r['signed'] is not None:
            ws.cell(rr, 5).font = red if r['signed'] < 0 else green
        if r['status'] in REVIEW_FLAGS:
            ws.cell(rr, 12).font = orange
    ws.auto_filter.ref = f"A1:L{ws.max_row}"
    rv = wb.create_sheet('Review')
    rv.append(HEADERS); _style_header(rv)
    for r in rec:
        if r['status'] in REVIEW_FLAGS:
            rv.append(_row_for(r)); rr = rv.max_row
            rv.cell(rr, 5).number_format = '#,##0.00'; rv.cell(rr, 6).number_format = '#,##0.00'
    rv.auto_filter.ref = f"A1:L{rv.max_row}"
    wb.save(path)


# ------------------------------------------------------------------------- main
def main(argv):
    args = [a for a in argv if a != '--max-balance']
    max_bal = MAX_BAL
    if '--max-balance' in argv:
        idx = argv.index('--max-balance')
        max_bal = int(argv[idx + 1])
        args = [a for k, a in enumerate(argv) if k not in (idx, idx + 1)]
    if len(args) < 2:
        sys.exit(__doc__)
    out_xlsx, pdfs = args[0], args[1:]
    for t in ('pdftoppm', 'pdfinfo', 'tesseract'):
        require_tool(t)

    all_recs, seq = [], 0
    for pdf in pdfs:
        if not os.path.exists(pdf):
            sys.exit(f"not found: {pdf}")
        npages = page_count(pdf)
        print(f"{os.path.basename(pdf)}: {npages} pages")
        for p in range(1, npages + 1):
            seq += 1
            tag = f"{seq:03d}"
            with tempfile.TemporaryDirectory() as wd:
                tsv = ocr_page(pdf, p, wd)
                if tsv:
                    all_recs += extract_page(tsv, tag)
            print(f"\r  ocr {p}/{npages}", end='', flush=True)
        print()

    rec = reconcile(all_recs, max_bal)
    sanitize_dates(rec)
    tx = [r for r in rec if not r['bf']]
    counts = Counter(r['status'] for r in rec)
    ok = counts['OK']; fixed = sum(v for k, v in counts.items() if k.startswith('FIX'))
    rev = sum(1 for r in tx if r['status'] in REVIEW_FLAGS)
    print(f"transactions: {len(tx)}  OK: {ok}  auto-fixed: {fixed}  review: {rev} "
          f"({rev / max(1, len(tx)) * 100:.1f}%)")
    if tx:
        print(f"final balance: {tx[-1]['final_balance']:,.2f}")
    write_excel(rec, out_xlsx)
    print(f"wrote {out_xlsx}  (sheets: Statement, Review)")


if __name__ == '__main__':
    main(sys.argv[1:])
